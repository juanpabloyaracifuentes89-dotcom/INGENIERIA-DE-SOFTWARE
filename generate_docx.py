# generate_docx.py
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

def set_cell_background(cell, fill_hex):
    tcPr = cell._element.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._element.get_or_add_tcPr()
    tcMar = parse_xml(f'<w:tcMar {nsdecls("w")}><w:top w:w="{top}" w:type="dxa"/><w:bottom w:w="{bottom}" w:type="dxa"/><w:left w:w="{left}" w:type="dxa"/><w:right w:w="{right}" w:type="dxa"/></w:tcMar>')
    tcPr.append(tcMar)

def create_document():
    doc = docx.Document()

    # Configurar márgenes estándar (1 pulgada = 2.54 cm)
    for section in doc.sections:
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)

    # Colores corporativos
    COLOR_PRIMARY = RGBColor(30, 58, 138)    # Azul institucional #1E3A8A
    COLOR_SECONDARY = RGBColor(71, 85, 105) # Slate #475569
    COLOR_DARK = RGBColor(15, 23, 42)       # Dark #0F172A

    # ==========================================
    # 1. PORTADA
    # ==========================================
    p_inst = doc.add_paragraph()
    p_inst.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_inst = p_inst.add_run("INSTITUCIÓN UNIVERSITARIA DIGITAL DE ANTIOQUIA\nFACULTAD DE INGENIERÍA\nPROGRAMA DE INGENIERÍA DE SOFTWARE")
    r_inst.font.name = "Arial"
    r_inst.font.size = Pt(13)
    r_inst.font.bold = True
    r_inst.font.color.rgb = COLOR_PRIMARY

    for _ in range(3):
        doc.add_paragraph()

    p_tit = doc.add_paragraph()
    p_tit.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_tit = p_tit.add_run("SISTEMA DE INFORMACIÓN PARA EL PARQUEADERO\n\"AUTOS COLOMBIA\"")
    r_tit.font.name = "Arial"
    r_tit.font.size = Pt(20)
    r_tit.font.bold = True
    r_tit.font.color.rgb = COLOR_PRIMARY

    p_sub = doc.add_paragraph()
    p_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_sub = p_sub.add_run("Evidencia de Aprendizaje 1 (EA1)\nIteración 1: Gestión de la Entrada y Salida de Vehículos")
    r_sub.font.name = "Arial"
    r_sub.font.size = Pt(14)
    r_sub.font.italic = True
    r_sub.font.color.rgb = COLOR_SECONDARY

    for _ in range(5):
        doc.add_paragraph()

    p_aut = doc.add_paragraph()
    p_aut.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_aut = p_aut.add_run("ESTUDIANTE:\n[Nombre y Apellidos del Estudiante]\n(Ejemplo de entrega: López, Karla)\n\nDOCENTE / ASESOR:\n[Nombre del Docente Asesor de la Asignatura]\n\nASIGNATURA:\nIngeniería de Software - Semestre 6 / Bloque 2")
    r_aut.font.name = "Arial"
    r_aut.font.size = Pt(11)
    r_aut.font.bold = True
    r_aut.font.color.rgb = COLOR_DARK

    for _ in range(3):
        doc.add_paragraph()

    p_fec = doc.add_paragraph()
    p_fec.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_fec = p_fec.add_run("Medellín, Colombia\nOctubre de 2026")
    r_fec.font.name = "Arial"
    r_fec.font.size = Pt(10)
    r_fec.font.color.rgb = COLOR_SECONDARY

    doc.add_page_break()

    # Helper para encabezados
    def add_h1(text):
        h = doc.add_heading(level=1)
        r = h.add_run(text)
        r.font.name = "Arial"
        r.font.size = Pt(15)
        r.font.bold = True
        r.font.color.rgb = COLOR_PRIMARY
        return h

    def add_h2(text):
        h = doc.add_heading(level=2)
        r = h.add_run(text)
        r.font.name = "Arial"
        r.font.size = Pt(13)
        r.font.bold = True
        r.font.color.rgb = COLOR_SECONDARY
        return h

    def add_h3(text):
        h = doc.add_heading(level=3)
        r = h.add_run(text)
        r.font.name = "Arial"
        r.font.size = Pt(11)
        r.font.bold = True
        r.font.color.rgb = COLOR_DARK
        return h

    def add_p(text, bold=False, italic=False):
        p = doc.add_paragraph()
        r = p.add_run(text)
        r.font.name = "Arial"
        r.font.size = Pt(10.5)
        r.font.bold = bold
        r.font.italic = italic
        p.paragraph_format.line_spacing = 1.15
        p.paragraph_format.space_after = Pt(6)
        return p

    def add_box(title, text):
        tbl = doc.add_table(rows=1, cols=1)
        tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
        cell = tbl.cell(0, 0)
        set_cell_background(cell, "F8FAFC")
        set_cell_margins(cell, top=140, bottom=140, left=200, right=200)
        p = cell.paragraphs[0]
        r1 = p.add_run(f"📌 {title}\n")
        r1.font.name = "Arial"
        r1.font.bold = True
        r1.font.size = Pt(10.5)
        r1.font.color.rgb = COLOR_PRIMARY
        r2 = p.add_run(text)
        r2.font.name = "Courier New"
        r2.font.size = Pt(9.5)
        p.paragraph_format.line_spacing = 1.15
        doc.add_paragraph()

    # ==========================================
    # 2. INTRODUCCIÓN
    # ==========================================
    add_h1("1. INTRODUCCIÓN")
    add_h2("1.1 Contexto Organizacional y Planteamiento del Problema")
    add_p("El parqueadero 'Autos Colombia' es una organización dedicada al servicio integral de estacionamiento y custodia vehicular. Su modelo de negocio atiende dos segmentos prioritarios: usuarios con contratos de mensualidad fija prepagada y usuarios ocasionales que abonan según el tiempo real de permanencia.")
    add_p("En la actualidad, las operaciones diarias de entrada y salida se realizan mediante métodos manuales en planillas de papel, lo cual desencadena las siguientes dificultades críticas:")
    add_p("• Tiempos de espera excesivos en la barrera de entrada debido a la consulta manual de contratos vigentes.\n• Desconocimiento en tiempo real sobre la disponibilidad física de celdas por tipo de vehículo (autos, motos, camionetas).\n• Reclamaciones injustificadas y disputas legales por daños vehiculares preexistentes que no son registrados al ingreso.\n• Pérdidas económicas por errores aritméticos en el cálculo manual de fracciones y horas al momento de la salida.")

    add_h2("1.2 Justificación y Objetivos")
    add_p("La implementación de un sistema de información moderno, con arquitectura en capas y base de datos relacional, permite a Autos Colombia blindar sus operaciones, agilizar el flujo de vehículos en patio a menos de 15 segundos por transacción y asegurar una liquidación exacta y auditable.")
    add_p("Objetivo General: Desarrollar la primera iteración del software de gestión para el parqueadero 'Autos Colombia', sistematizando el ciclo de entrada, salida, control de mensualidades, celdas y registro de novedades.", bold=True)

    # ==========================================
    # 3. REQUERIMIENTOS FUNCIONALES Y NO FUNCIONALES
    # ==========================================
    add_h1("2. ENTREGABLE 1: REQUERIMIENTOS FUNCIONALES Y NO FUNCIONALES")
    add_h2("2.1 Requerimientos Funcionales (RF)")
    
    rf_data = [
        ("RF-01", "Validación de Mensualidad", "El sistema debe verificar de forma automática si la placa digitada cuenta con contrato de mensualidad vigente.", "Must Have"),
        ("RF-02", "Control Anti-Passback", "El sistema debe bloquear el registro de entrada si el vehículo ya figura con un registro activo dentro del parqueadero.", "Must Have"),
        ("RF-03", "Registro de Entrada", "Permite capturar placa, tipo de automotor, celda asignada, hora exacta y operador de turno, marcando la celda como ocupada.", "Must Have"),
        ("RF-04", "Asignación Asistida de Celdas", "Filtra y muestra en tiempo real únicamente celdas disponibles compatibles con el tipo de vehículo a ingresar.", "Must Have"),
        ("RF-05", "Registro de Novedades e Incidentes", "Permite tipificar averías previas (rayones, golpes, accesorios) con descripción, ubicación y gravedad (Leve, Moderada, Grave).", "Should Have"),
        ("RF-06", "Mapa Gráfico de Celdas", "Provee una cuadrícula con estados visuales (Verde: Libre, Rojo: Ocupada, Amarillo: Mantenimiento) y placa actual.", "Must Have"),
        ("RF-07", "Registro de Salida y Liberación", "Permite seleccionar un automotor en patio, calcular su permanencia y liberar la celda a estado disponible.", "Must Have"),
        ("RF-08", "Liquidación Automática de Tarifas", "Calcula el valor a pagar ($0 COP para mensualidades vigentes o cobro fraccionado/hora para visitantes).", "Must Have"),
        ("RF-09", "Emisión de Tickets y Facturas", "Genera comprobantes impresos/digitales de entrada y salida con soporte de póliza de responsabilidad civil.", "Should Have"),
        ("RF-10", "Auditoría de Novedades", "Permite consultar el histórico consolidado de incidentes reportados para fines legales y administrativos.", "Could Have")
    ]

    tbl_rf = doc.add_table(rows=len(rf_data)+1, cols=4)
    tbl_rf.alignment = WD_TABLE_ALIGNMENT.CENTER
    headers_rf = ["Código", "Nombre Requerimiento", "Descripción Detallada", "Prioridad"]
    for j, h in enumerate(headers_rf):
        cell = tbl_rf.cell(0, j)
        set_cell_background(cell, "1E3A8A")
        set_cell_margins(cell, top=100, bottom=100, left=120, right=120)
        p = cell.paragraphs[0]
        r = p.add_run(h)
        r.font.name = "Arial"
        r.font.bold = True
        r.font.size = Pt(9.5)
        r.font.color.rgb = RGBColor(255, 255, 255)

    for i, row in enumerate(rf_data):
        for j, val in enumerate(row):
            cell = tbl_rf.cell(i+1, j)
            if i % 2 == 1:
                set_cell_background(cell, "F1F5F9")
            set_cell_margins(cell, top=80, bottom=80, left=100, right=100)
            p = cell.paragraphs[0]
            r = p.add_run(val)
            r.font.name = "Arial"
            r.font.size = Pt(9)
            if j == 0:
                r.font.bold = True

    doc.add_paragraph()

    add_h2("2.2 Requerimientos No Funcionales (RNF - FURPS+)")
    rnf_data = [
        ("RNF-01", "Rendimiento", "La validación de placa contra contratos de mensualidad debe responder en menos de 800 ms."),
        ("RNF-02", "Concurrencia", "El backend debe soportar al menos 50 solicitudes simultáneas sin pérdida de rendimiento."),
        ("RNF-03", "Seguridad ACID", "Toda asignación de celda e ingreso debe ejecutarse bajo transacciones atómicas con llaves foráneas."),
        ("RNF-04", "Trazabilidad", "Todo registro de entrada, salida o daño almacena inmutablemente la marca de tiempo y el operador."),
        ("RNF-05", "Usabilidad (UX)", "El registro de entrada no debe requerir más de 3 clics ni más de 15 segundos en caseta."),
        ("RNF-06", "Responsividad", "La interfaz debe operar con perfecta visualización en pantallas de escritorio y tablets de 10 pulgadas."),
        ("RNF-07", "Confiabilidad", "Disponibilidad del 99.5% con persistencia en disco local protegida contra fallos eléctricos."),
        ("RNF-08", "Mantenibilidad", "Arquitectura en capas, API REST estandarizada y base de datos SQL normalizada.")
    ]

    tbl_rnf = doc.add_table(rows=len(rnf_data)+1, cols=3)
    tbl_rnf.alignment = WD_TABLE_ALIGNMENT.CENTER
    headers_rnf = ["Código", "Categoría FURPS+", "Especificación Técnica"]
    for j, h in enumerate(headers_rnf):
        cell = tbl_rnf.cell(0, j)
        set_cell_background(cell, "334155")
        set_cell_margins(cell, top=100, bottom=100, left=120, right=120)
        p = cell.paragraphs[0]
        r = p.add_run(h)
        r.font.name = "Arial"
        r.font.bold = True
        r.font.size = Pt(9.5)
        r.font.color.rgb = RGBColor(255, 255, 255)

    for i, row in enumerate(rnf_data):
        for j, val in enumerate(row):
            cell = tbl_rnf.cell(i+1, j)
            if i % 2 == 1:
                set_cell_background(cell, "F8FAFC")
            set_cell_margins(cell, top=80, bottom=80, left=100, right=100)
            p = cell.paragraphs[0]
            r = p.add_run(val)
            r.font.name = "Arial"
            r.font.size = Pt(9)
            if j == 0:
                r.font.bold = True

    doc.add_paragraph()

    # ==========================================
    # 4. HISTORIAS DE USUARIO
    # ==========================================
    add_h1("3. ENTREGABLE 2: HISTORIAS DE USUARIO ASOCIADAS")
    add_p("A continuación se presentan las Historias de Usuario bajo el estándar ágil INVEST, incorporando criterios de aceptación estructurados en lenguaje Given-When-Then (Gherkin):")

    add_box("HU-01: Registro y Validación de Mensualidad",
            "Como: Operador de patio del parqueadero\n"
            "Quiero: Digitar la placa de un vehículo para verificar si cuenta con mensualidad activa\n"
            "Para: Conceder el ingreso expedito sin costo y asignarle su celda correspondiente.\n\n"
            "[Criterio de Aceptación 1: Mensualidad Vigente]\n"
            "Dado que se digita la placa 'ABC123' asociada a un contrato con vencimiento futuro\n"
            "Cuando se verifica la placa en el sistema\n"
            "Entonces se muestra el banner verde de confirmación con el nombre del cliente y valor $0 COP.\n\n"
            "[Criterio de Aceptación 2: Control de Duplicidad (Anti-Passback)]\n"
            "Dado que el vehículo 'XYZ789' ya está registrado activo en patio\n"
            "Cuando se intente ingresar nuevamente dicha placa\n"
            "Entonces el sistema bloquea la acción informando que ya se encuentra en la celda M-01.")

    add_box("HU-02: Asignación Visual de Celda Compatible",
            "Como: Operador de patio\n"
            "Quiero: Visualizar exclusivamente celdas disponibles acordes al tipo de vehículo\n"
            "Para: Asignar un espacio físico adecuado y optimizar la circulación del patio.\n\n"
            "[Criterio de Aceptación]\n"
            "Dado que ingresa una motocicleta\n"
            "Cuando el operador selecciona tipo 'Motocicleta'\n"
            "Entonces el desplegable solo muestra celdas de moto disponibles (ej. M-01, M-02).")

    add_box("HU-03: Levantamiento de Novedades e Inventario al Ingreso",
            "Como: Operador de recepción\n"
            "Quiero: Documentar rayones, abolladuras o piezas averiadas en el vehículo al entrar\n"
            "Para: Proteger a la empresa de reclamaciones indebidas al momento de la salida.\n\n"
            "[Criterio de Aceptación]\n"
            "Dado que un vehículo ingresa con un rayón en la puerta trasera\n"
            "Cuando se activa la casilla de novedades y se detallan tipo, gravedad y descripción\n"
            "Entonces el registro de acceso guarda la novedad enlazada para auditoría.")

    add_box("HU-05: Liquidación Automática y Registro de Salida",
            "Como: Operador de caseta\n"
            "Quiero: Dar salida a un vehículo calculando automáticamente el tiempo y valor a pagar\n"
            "Para: Liberar la celda inmediatamente en el sistema y cobrar la tarifa exacta.\n\n"
            "[Criterio de Aceptación]\n"
            "Dado que sale el vehículo 'QWE987' que permaneció 120 minutos (2 horas)\n"
            "Cuando se confirma la salida\n"
            "Entonces se cobra $9.000 COP (2 x $4.500), la celda pasa a LIBRE y se emite la factura.")

    # ==========================================
    # 5. ARQUITECTURA DEL SISTEMA
    # ==========================================
    add_h1("4. ENTREGABLE 3: ARQUITECTURA DEL SISTEMA EN STARUML")
    add_p("El sistema está concebido bajo una Arquitectura en Tres Capas (Three-Tier Architecture), la cual garantiza desacoplamiento, mantenibilidad y alta cohesión:")
    add_p("1. Capa de Presentación (Frontend): Desarrollada en HTML5, CSS3 moderno y JavaScript Vanilla ES6. Ofrece al usuario una interfaz reactiva y visual para la gestión de celdas en tiempo real.")
    add_p("2. Capa de Aplicación / Lógica de Negocio (Backend Express): Implementada en Node.js mediante el framework Express.js. Administra las rutas REST, ejecuta las validaciones de anti-passback, evalúa vigencias de mensualidad y calcula las liquidaciones tarifarias.")
    add_p("3. Capa de Persistencia y Datos: Base de datos relacional SQLite embebida de alto desempeño (`node:sqlite`), con claves foráneas e integridad referencial garantizada bajo ACID.")
    add_p("Nota de Entrega StarUML: El diagrama de componentes y despliegue de arquitectura se encuentra completamente especificado en el archivo 'diagramas_staruml/arquitectura_sistema.puml', listo para visualización o importación en StarUML.", italic=True)

    # ==========================================
    # 6. MOCKUPS EN FIGMA
    # ==========================================
    add_h1("5. ENTREGABLE 4: MOCKUPS DE LAS HISTORIAS DE USUARIO EN FIGMA")
    add_p("El diseño de interfaz (UI/UX) se fundamentó en un Sistema de Diseño con contraste optimizado para entornos de caseta:")
    add_p("• Paleta Cromática: Azul Primario (#1E3A8A), Verde Éxito (#10B981) para disponibles, Rojo Alerta (#EF4444) para ocupadas, y Amarillo (#FEF08A) para placas vehiculares colombianas.")
    add_p("• Pantallas Diseñadas en el Flujo:\n  1. Dashboard de Patio: Cuadrícula interactiva con las 12 celdas de patio y contadores KPI en vivo.\n  2. Modal de Entrada: Formulario con verificación de placa, banner de mensualidad y acordeón de novedades.\n  3. Modal de Salida: Resumen de permanencia, horas liquidadas y valor total.\n  4. Ticket de Comprobante: Formato para impresora térmica con datos fiscales y póliza civil.\n  5. Libro de Novedades: Tabla histórica de incidentes tipificados.")
    add_p("Nota de Figma: Los wireframes y lineamientos para replicar o enlazar el prototipo interactivo en Figma se encuentran detallados en 'figma_mockups/guia_mockups_figma.md'.", italic=True)

    # ==========================================
    # 7. MODELOS DE DATOS (MER, MR Y FÍSICO)
    # ==========================================
    add_h1("6. ENTREGABLE 5: MER, MR Y MODELO FÍSICO")
    add_h2("6.1 Modelo Entidad-Relación (MER Conceptual)")
    add_p("Entidades Principales y Cardinalidades:\n• TIPOS_VEHICULO (1) a VEHICULOS (N): Un tipo clasifica a múltiples vehículos.\n• CLIENTES (1) a MENSUALIDADES (N): Un cliente puede contratar una o varias mensualidades.\n• VEHICULOS (1) a REGISTROS_ACCESO (N): Un vehículo acumula un historial de entradas y salidas.\n• CELDAS (1) a REGISTROS_ACCESO (N): Una celda aloja múltiples accesos a lo largo del tiempo.\n• REGISTROS_ACCESO (1) a NOVEDADES (N): Un registro de acceso documenta cero o más novedades.")

    add_h2("6.2 Modelo Relacional Normalizado (3FN)")
    add_p("• tipos_vehiculo (id [PK], nombre [UQ], tarifa_hora, tarifa_dia, tarifa_mes, icono)\n• clientes (id [PK], documento [UQ], nombres, apellidos, telefono, email, estado, fecha_creacion)\n• vehiculos (placa [PK], id_tipo_vehiculo [FK], marca, modelo, color, id_cliente [FK])\n• mensualidades (id [PK], id_cliente [FK], placa [FK], fecha_inicio, fecha_fin, valor_pagado, estado)\n• celdas (codigo [PK], id_tipo_vehiculo [FK], piso, estado, placa_actual [FK])\n• registros_acceso (id [PK], placa [FK], codigo_celda [FK], fecha_hora_entrada, fecha_hora_salida, tipo_servicio, minutos_transcurridos, total_pagar, estado, observaciones, operador_entrada, operador_salida)\n• novedades (id [PK], id_registro_acceso [FK], placa [FK], fecha_hora, tipo_novedad, nivel_gravedad, descripcion, ubicacion)")

    add_h2("6.3 Modelo Físico SQL")
    add_p("El script completo de creación de base de datos DDL con restricciones e índices optimizados se ubica en el archivo ejecutable 'database/schema.sql'.")

    # ==========================================
    # 8. DIAGRAMAS UML (CASOS DE USO, CLASES, PAQUETES)
    # ==========================================
    add_h1("7. ENTREGABLES 6, 7 Y 8: DIAGRAMAS UML EN STARUML")
    add_h2("7.1 Diagrama de Casos de Uso (Punto 6)")
    add_p("Modela las interacciones del Operador de Patio y el Administrador con el sistema. Incluye las relaciones obligatorias <<include>> para la validación de mensualidades y liquidación de salidas, y <<extend>> para el registro opcional de novedades al ingreso.")
    add_p("Archivo fuente: 'diagramas_staruml/casos_de_uso.puml'.", italic=True)

    add_h2("7.2 Diagrama de Clases del Dominio (Punto 7)")
    add_p("Estructura las entidades del software con tipado estricto, encapsulamiento y métodos de negocio: Vehiculo, TipoVehiculo, Cliente, Mensualidad, Celda, RegistroAcceso y Novedad, complementadas con enumeraciones para estados y gravedades.")
    add_p("Archivo fuente: 'diagramas_staruml/diagrama_clases.puml'.", italic=True)

    add_h2("7.3 Diagrama de Paquetes Modular (Punto 8)")
    add_p("Organiza la solución en los paquetes co.edu.autoscolombia: presentation, controllers, services, models, repositories y config, garantizando un flujo unidireccional de dependencias.")
    add_p("Archivo fuente: 'diagramas_staruml/diagrama_paquetes.puml'.", italic=True)

    # ==========================================
    # 9. CÓDIGO FUENTE Y GITHUB
    # ==========================================
    add_h1("8. ENTREGABLE 9: CÓDIGO FUENTE CORRIENDO Y REPOSITORIO GITHUB")
    add_p("El sistema fue implementado íntegramente y se encuentra 100% operativo bajo Node.js y SQLite:")
    add_p("• Servidor Web: Express.js en el puerto 3000 con API REST documentada.\n• Base de Datos: SQLite con auto-migración y carga de datos semilla (12 celdas, tarifas oficiales y clientes de prueba).\n• Frontend: Dashboard interactivo con mapa de celdas en vivo, registro de entradas con modal asistido, registro de salidas con cálculo en vivo de cobro y tickets imprimibles.")
    add_p("Instrucciones de Ejecución Local:\n1. Clonar o abrir la carpeta del proyecto en terminal.\n2. Ejecutar 'npm install' y luego 'npm start'.\n3. Acceder en el navegador a: http://localhost:3000")
    add_p("Instrucciones para GitHub:\nEl repositorio local Git ya ha sido inicializado y comiteado en la rama 'main'. Para publicarlo en GitHub, basta con ejecutar:\n  git remote add origin https://github.com/TU_USUARIO/autos-colombia-parqueadero.git\n  git push -u origin main")

    # ==========================================
    # 10. CONCLUSIONES
    # ==========================================
    add_h1("9. CONCLUSIONES")
    add_p("1. La aplicación rigurosa de los conceptos de ingeniería de software en la Iteración 1 permitió diseñar una solución integral, robusta y alineada con los desafíos reales del parqueadero 'Autos Colombia'.")
    add_p("2. El módulo de levantamiento digital de novedades mitiga de manera efectiva el riesgo legal ante reclamaciones por daños preexistentes, otorgando trazabilidad total con marca de tiempo y operador responsable.")
    add_p("3. La arquitectura en tres capas y la base de datos relacional normalizada en 3FN proporcionan una base sólida, desacoplada y escalable para la integración futura de las iteraciones 2 y 3 (facturación electrónica y gestión de usuarios avanzados).")

    # Guardar documento
    doc_path = r"c:\Users\USER\Desktop\Universidad\SEMESTRE 6\BLOQUE 2\INGENIERIA DE SOFTWARE\Autos_Colombia_EA1_Documento.docx"
    doc.save(doc_path)
    print(f"Documento Word generado exitosamente en: {doc_path}")

if __name__ == "__main__":
    create_document()
